from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "Jetson_Orin_Nano_Super_Vision_VLA_가이드북.docx"

CHAPTERS = [
    ROOT / "docs" / "01_하드웨어와_전원.md",
    ROOT / "docs" / "02_핀과_커넥터.md",
    ROOT / "docs" / "03_JetPack_BSP_설치.md",
    ROOT / "docs" / "04_커스텀_BSP_브링업.md",
    ROOT / "docs" / "05_카메라와_Vision.md",
    ROOT / "docs" / "06_AI_가속과_최적화.md",
    ROOT / "docs" / "07_ROS2_Isaac_ROS.md",
    ROOT / "docs" / "08_Vision_VLA_로드맵.md",
    ROOT / "docs" / "09_운영_안전_문제해결.md",
    ROOT / "references" / "공식문서_목록.md",
    ROOT / "references" / "커버리지_매트릭스.md",
]

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
NAVY = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x66, 0x6D, 0x75)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BODY_FONT = "Calibri"
KOREAN_FONT = "Malgun Gothic"
MONO_FONT = "Consolas"
TOTAL_DXA = 9360


def set_run_font(run, size=None, bold=None, italic=None, color=None, mono=False):
    name = MONO_FONT if mono else BODY_FONT
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), KOREAN_FONT if not mono else MONO_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_style_font(style, size, color=None, bold=None):
    style.font.name = BODY_FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, 9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_hyperlink(paragraph, label, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), BODY_FONT)
    rfonts.set(qn("w:hAnsi"), BODY_FONT)
    rfonts.set(qn("w:eastAsia"), KOREAN_FONT)
    rpr.extend([rfonts, color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([rpr, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")
CODE_RE = re.compile(r"`([^`]+)`")


def add_inline(paragraph, text, bold=False):
    cursor = 0
    tokens = []
    for match in LINK_RE.finditer(text):
        if match.start() > cursor:
            tokens.append(("text", text[cursor:match.start()], None))
        tokens.append(("link", match.group(1), match.group(2)))
        cursor = match.end()
    if cursor < len(text):
        tokens.append(("text", text[cursor:], None))
    if not tokens:
        tokens = [("text", text, None)]

    for kind, value, url in tokens:
        if kind == "link":
            add_hyperlink(paragraph, value, url)
            continue
        pos = 0
        for code in CODE_RE.finditer(value):
            if code.start() > pos:
                run = paragraph.add_run(value[pos:code.start()])
                set_run_font(run, 11, bold=bold)
            run = paragraph.add_run(code.group(1))
            set_run_font(run, 9.5, bold=bold, color=DARK_BLUE, mono=True)
            pos = code.end()
        if pos < len(value):
            run = paragraph.add_run(value[pos:])
            set_run_font(run, 11, bold=bold)


def add_numbering(document):
    numbering = document.part.numbering_part.element

    def make(abstract_id, num_id, fmt, text, font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.extend([tabs, ind, spacing])
        lvl.extend([start, num_fmt, lvl_text, suff, ppr])
        if font:
            rpr = OxmlElement("w:rPr")
            rfonts = OxmlElement("w:rFonts")
            rfonts.set(qn("w:ascii"), font)
            rfonts.set(qn("w:hAnsi"), font)
            rpr.append(rfonts)
            lvl.append(rpr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_num_id = OxmlElement("w:abstractNumId")
        abstract_num_id.set(qn("w:val"), str(abstract_id))
        num.append(abstract_num_id)
        numbering.append(num)

    make(910, 910, "bullet", "•", "Symbol")
    make(911, 911, "decimal", "%1.")
    return 910, 911


def apply_num(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, numid])
    ppr.append(num_pr)


def parse_table(lines, start):
    rows = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        cells = [c.strip() for c in lines[idx].strip().strip("|").split("|")]
        rows.append(cells)
        idx += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        return [rows[0]] + rows[2:], idx
    return None, start


def table_widths(headers):
    n = len(headers)
    if n == 2:
        if all(re.search(r"핀|Pin", h, re.I) for h in headers[::2] if headers):
            return [1600, 7760]
        return [2700, 6660]
    if n == 4 and headers[0] in {"핀", "Pin"}:
        return [900, 3780, 900, 3780]
    if n == 5:
        return [1500, 2100, 2300, 900, 2560]
    base = TOTAL_DXA // n
    widths = [base] * n
    widths[-1] += TOTAL_DXA - sum(widths)
    return widths


def add_markdown_table(doc, rows):
    max_cols = max(len(r) for r in rows)
    rows = [r + [""] * (max_cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = "Table Grid"
    # Mark the first row as a semantic/repeating header for screen readers
    # and for long tables that flow across Word pages.
    header_props = table.rows[0]._tr.get_or_add_trPr()
    header_props.append(OxmlElement("w:tblHeader"))
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value, bold=(r_idx == 0))
            for run in p.runs:
                set_run_font(run, 9 if max_cols >= 4 else 9.5, bold=(r_idx == 0))
            if r_idx == 0:
                shade_cell(cell, LIGHT_BLUE)
    set_table_geometry(table, table_widths(rows[0]))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(doc, code):
    p = doc.add_paragraph(style="Code Block")
    for line in code.rstrip().splitlines():
        run = p.add_run(line)
        set_run_font(run, 8.5, mono=True, color=NAVY)
        run.add_break()


def add_markdown(doc, path, bullet_id, decimal_id, chapter_break=True):
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    if chapter_break:
        doc.add_page_break()
    idx = 0
    in_code = False
    code_lines = []
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            p = doc.add_paragraph()
            add_inline(p, " ".join(s.strip() for s in paragraph_buffer))
            paragraph_buffer = []

    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                add_code_block(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue
        if in_code:
            code_lines.append(line)
            idx += 1
            continue
        if not stripped:
            flush_paragraph()
            idx += 1
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            rows, next_idx = parse_table(lines, idx)
            if rows:
                add_markdown_table(doc, rows)
                idx = next_idx
                continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = min(len(heading.group(1)), 3)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, heading.group(2), bold=True)
            idx += 1
            continue
        if re.match(r"^-\s+", stripped):
            flush_paragraph()
            p = doc.add_paragraph()
            apply_num(p, bullet_id)
            add_inline(p, re.sub(r"^-\s+", "", stripped))
            idx += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            flush_paragraph()
            p = doc.add_paragraph()
            apply_num(p, decimal_id)
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
            idx += 1
            continue
        paragraph_buffer.append(stripped)
        idx += 1
    flush_paragraph()


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        set_style_font(style, size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(code, 8.5, color=NAVY)
    code.font.name = MONO_FONT
    code._element.rPr.rFonts.set(qn("w:ascii"), MONO_FONT)
    code._element.rPr.rFonts.set(qn("w:hAnsi"), MONO_FONT)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.08)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0
    ppr = code._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    ppr.append(shd)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("JETSON ORIN NANO SUPER  |  VISION VLA FIELD GUIDE")
    set_run_font(run, 8.5, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("JETSON ORIN NANO SUPER")
    set_run_font(run, 12, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("임베디드 Vision VLA\n개발 가이드북")
    set_run_font(run, 28, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    run = p.add_run("하드웨어 · 핀 · JetPack/BSP · 카메라 · AI 가속 · ROS 2 · 안전한 행동 배포")
    set_run_font(run, 13, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("기준: JetPack 7.2.1 / Jetson Linux 39.2.1 / ROS 2 Jazzy")
    set_run_font(run, 10.5, bold=True, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("문서 기준일 2026-08-29")
    set_run_font(run, 10, color=MUTED)

    doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    add_inline(p, "사용 안내", bold=True)
    add_inline(doc.add_paragraph(), "이 책은 공식 NVIDIA/ROS 문서를 한국어로 재서술한 실무 해설서다. 전기 절대 정격, 핀 drive, 플래시 명령, Secure Boot 키 작업은 설치한 릴리스와 동일한 공식 원문을 최종 기준으로 확인한다.")
    p = doc.add_paragraph(style="Heading 2")
    add_inline(p, "권장 읽기 경로", bold=True)
    for item in (
        "처음 설치: 1장 → 2장 → 3장 → 5장 → 6장",
        "커스텀 캐리어보드: 2장 → 4장 → 9장",
        "Vision VLA: 5장 → 6장 → 7장 → 8장 → 9장",
    ):
        para = doc.add_paragraph()
        apply_num(para, 910)
        add_inline(para, item)

    p = doc.add_paragraph(style="Heading 2")
    add_inline(p, "목차", bold=True)
    for path in CHAPTERS:
        title = path.read_text(encoding="utf-8").splitlines()[0].lstrip("# ")
        para = doc.add_paragraph()
        apply_num(para, 911)
        add_inline(para, title)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    bullet_id, decimal_id = add_numbering(doc)
    add_cover(doc)
    for chapter in CHAPTERS:
        add_markdown(doc, chapter, bullet_id, decimal_id, chapter_break=True)
    props = doc.core_properties
    props.title = "Jetson Orin Nano Super 임베디드 Vision VLA 개발 가이드북"
    props.subject = "핀, 전원, JetPack/BSP, 카메라, TensorRT, ROS 2, Isaac ROS, Vision VLA"
    props.author = "RoboticsStudy"
    props.keywords = "Jetson, Orin Nano Super, BSP, Vision, VLA, ROS 2, Isaac ROS"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
