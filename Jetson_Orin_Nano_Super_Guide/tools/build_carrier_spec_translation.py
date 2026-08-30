from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, Spacer

import build_guidebook as word
import build_pdf as pdf


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "translations" / "Jetson_Orin_Nano_DevKit_Carrier_Board_Spec_v1.3_KO.md"
DOCX_OUT = ROOT / "output" / "Jetson_Orin_Nano_Carrier_Board_Spec_v1.3_KO.docx"
PDF_OUT = ROOT / "output" / "Jetson_Orin_Nano_Carrier_Board_Spec_v1.3_KO.pdf"


def add_word_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    word.set_run_font(p.add_run("NVIDIA JETSON ORIN NANO"), 12, bold=True, color=word.BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    word.set_run_font(p.add_run("개발자 키트 캐리어보드\n사양서 v1.3"), 27, bold=True, color=word.NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    word.set_run_font(p.add_run("한국어 기술 번역 · 핀/전원 설계 해설"), 14, color=word.DARK_BLUE)
    for text in (
        "원문: SP-11324-001_v1.3 | December 2024",
        "번역 기준일: 2026-08-30",
        "기술 이해용 번역 - 전기 설계는 NVIDIA 최신 원문 우선",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        word.set_run_font(p.add_run(text), 10.5, bold=text.startswith("원문:"), color=word.MUTED)
    doc.add_page_break()


def build_docx():
    doc = Document()
    word.configure_document(doc)
    header = doc.sections[0].header.paragraphs[0]
    header.text = ""
    word.set_run_font(header.add_run("JETSON ORIN NANO  |  CARRIER BOARD SPEC v1.3 KO"), 8.5, bold=True, color=word.MUTED)
    bullet_id, decimal_id = word.add_numbering(doc)
    add_word_cover(doc)
    word.add_markdown(doc, SOURCE, bullet_id, decimal_id, chapter_break=False)
    props = doc.core_properties
    props.title = "Jetson Orin Nano 개발자 키트 캐리어보드 사양서 v1.3 한국어 번역"
    props.subject = "커넥터, 핀, 전원, 기구 사양 한국어 기술 번역"
    props.author = "RoboticsStudy"
    props.keywords = "Jetson Orin Nano, carrier board, pinout, power, Korean translation"
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)


class TranslationDocTemplate(pdf.GuideDocTemplate):
    def draw_page(self, canvas, doc):
        canvas.saveState()
        canvas.setFont("Malgun-Bold", 7.5)
        canvas.setFillColor(pdf.MUTED)
        canvas.drawString(doc.leftMargin, letter[1] - 0.48 * inch, "JETSON ORIN NANO  |  CARRIER BOARD SPEC v1.3 KO")
        canvas.setFont("Malgun", 8)
        canvas.drawRightString(letter[0] - doc.rightMargin, 0.47 * inch, f"Page {doc.page}")
        canvas.restoreState()


def pdf_cover(st):
    title = ParagraphStyle("TranslationCoverTitle", parent=st["h1"], fontSize=26, leading=34, alignment=TA_CENTER, textColor=pdf.NAVY, spaceAfter=14)
    kicker = ParagraphStyle("TranslationCoverKicker", parent=st["small"], fontName="Malgun-Bold", fontSize=11, alignment=TA_CENTER, textColor=pdf.BLUE, spaceAfter=16)
    subtitle = ParagraphStyle("TranslationCoverSubtitle", parent=st["body"], fontSize=13, leading=19, alignment=TA_CENTER, textColor=pdf.DARK_BLUE, spaceAfter=30)
    meta = ParagraphStyle("TranslationCoverMeta", parent=st["small"], alignment=TA_CENTER, spaceAfter=5)
    return [
        Spacer(1, 1.35 * inch),
        Paragraph("NVIDIA JETSON ORIN NANO", kicker),
        Paragraph("개발자 키트 캐리어보드<br/>사양서 v1.3", title),
        Paragraph("한국어 기술 번역 · 핀/전원 설계 해설", subtitle),
        Paragraph("원문: SP-11324-001_v1.3 | December 2024", meta),
        Paragraph("번역 기준일: 2026-08-30", meta),
        Paragraph("기술 이해용 번역 - 전기 설계는 NVIDIA 최신 원문 우선", meta),
        PageBreak(),
    ]


def build_pdf():
    pdf.register_fonts()
    st = pdf.styles()
    doc = TranslationDocTemplate(
        str(PDF_OUT),
        pagesize=letter,
        leftMargin=0.78 * inch,
        rightMargin=0.78 * inch,
        topMargin=0.78 * inch,
        bottomMargin=0.72 * inch,
        title="Jetson Orin Nano 캐리어보드 사양서 v1.3 한국어 번역",
        author="RoboticsStudy",
        subject="핀, 커넥터, 전원 및 기구 사양 번역",
    )
    story = pdf_cover(st)
    story.extend(pdf.markdown_story(SOURCE, st, doc.width))
    PDF_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX_OUT)
    print(PDF_OUT)
